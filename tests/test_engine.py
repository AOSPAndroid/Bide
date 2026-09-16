import base64
import io
import zipfile
from pathlib import Path

import pymupdf as fitz
import pytest
from fastapi.testclient import TestClient
from PIL import Image
from docx import Document
from server import app, SOURCES, libreoffice

client = TestClient(app)


@pytest.fixture()
def pdf():
    doc = fitz.open()
    for text in ['Original first page', 'Original second page']:
        page = doc.new_page(width=595, height=842)
        page.insert_text((50, 70), text, fontsize=22)
        page.draw_rect(fitz.Rect(50, 100, 250, 200), fill=(0.2, 0.7, 0.5))
    doc[0].insert_link({'kind':fitz.LINK_URI, 'from':fitz.Rect(50,50,280,80), 'uri':'https://example.com'})
    return doc.tobytes()


def register(pdf):
    response = client.post('/api/import', files={'file':('sample.pdf',pdf,'application/pdf')})
    assert response.status_code == 200, response.text
    return response.json()


def payload(source, svg='', page_index=0):
    return {'pages':[{'width':595,'height':842,'source':source['id'],'index':page_index,'svg':svg}], 'title':'Test document'}


def test_import_render_and_roundtrip(pdf):
    source = register(pdf)
    assert len(source['pages']) == 2
    response = client.get(f"/api/sources/{source['id']}/pages/0?scale=0.4")
    assert response.status_code == 200
    assert Image.open(io.BytesIO(response.content)).size == (238,337)
    svg = '''<svg xmlns="http://www.w3.org/2000/svg" width="595" height="842" viewBox="0 0 595 842"><g transform="translate(90 280)"><text font-family="Arial" font-size="25" fill="#173e33">Added editable text</text></g><rect x="50" y="400" width="100" height="40" fill="#ff0000"/></svg>'''
    response = client.post('/api/export',json=payload(source,svg))
    assert response.status_code == 200, response.text
    with fitz.open(stream=response.content,filetype='pdf') as exported:
        assert len(exported) == 1
        assert 'Original first page' in exported[0].get_text()
        assert 'Added editable text' in exported[0].get_text()
        assert exported[0].get_links()[0]['uri'] == 'https://example.com'
        assert len(exported[0].get_drawings()) >= 2
        assert exported[0].rect == fitz.Rect(0,0,595,842)
        out = Path('tmp/pdfs'); out.mkdir(parents=True,exist_ok=True)
        (out/'roundtrip.pdf').write_bytes(response.content)
        (out/'source.pdf').write_bytes(pdf)


def test_merge_reorder_and_split(pdf):
    source = register(pdf)
    data = payload(source)
    data['pages'] = [payload(source,page_index=i)['pages'][0] for i in [1,0,1]]
    response = client.post('/api/export',json=data)
    with fitz.open(stream=response.content,filetype='pdf') as merged:
        assert len(merged) == 3
        assert 'second' in merged[0].get_text()
        assert 'first' in merged[1].get_text()
    response = client.post('/api/export',json={**data,'format':'split','range':'1,3'})
    assert response.status_code == 200, response.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert len(archive.namelist()) == 2
        for name in archive.namelist():
            with fitz.open(stream=archive.read(name),filetype='pdf') as doc:
                assert len(doc) == 1 and 'second' in doc[0].get_text()


@pytest.mark.parametrize('format',['png','jpg','svg','docx','txt'])
def test_conversions(pdf,format):
    response = client.post('/api/export',json={**payload(register(pdf)),'format':format})
    assert response.status_code == 200, response.text
    if format in ['png','jpg','svg']:
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            assert len(archive.namelist()) == 1
            data=archive.read(archive.namelist()[0])
            if format!='svg':
                assert Image.open(io.BytesIO(data)).size == (1190,1684)
    elif format=='docx':
        assert 'Original first page' in '\n'.join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)
    else:
        assert 'Original first page' in response.text


def test_project_restore(pdf):
    source=register(pdf)
    SOURCES.pop(source['id'])
    response=client.post('/api/restore',json={source['id']:source['data']})
    assert response.status_code == 200
    assert client.post('/api/export',json=payload(source)).status_code == 200
    assert client.post('/api/restore',json={source['id']:base64.b64encode(b'bad').decode()}).status_code == 422


def test_compression_preserves_text():
    image=Image.effect_noise((2400,2400),60).convert('RGB')
    buf=io.BytesIO(); image.save(buf,'PNG')
    doc=fitz.open(); page=doc.new_page(width=595,height=842)
    page.insert_image(fitz.Rect(0,100,595,695),stream=buf.getvalue())
    page.insert_text((50,50),'Keep this text',fontsize=18)
    source=register(doc.tobytes())
    original=client.post('/api/export',json=payload(source))
    small=client.post('/api/export',json={**payload(source),'compression':'small'})
    assert small.status_code == 200, small.text
    assert len(small.content) < len(original.content)/2
    with fitz.open(stream=small.content,filetype='pdf') as result:
        assert 'Keep this text' in result[0].get_text()


def test_rotated_source_and_forms(pdf):
    doc=fitz.open(stream=pdf,filetype='pdf'); doc[0].set_rotation(90)
    widget=fitz.Widget(); widget.field_name='Test field'; widget.field_type=fitz.PDF_WIDGET_TYPE_TEXT
    widget.rect=fitz.Rect(50,300,200,330); widget.field_value='preserved'
    doc[0].add_widget(widget)
    source=register(doc.tobytes())
    data={'pages':[{'width':842,'height':595,'source':source['id'],'index':0,'svg':''}]}
    response=client.post('/api/export',json=data)
    assert response.status_code==200, response.text
    with fitz.open(stream=response.content,filetype='pdf') as result:
        assert result[0].rect == fitz.Rect(0,0,842,595)
        assert 'Original first page' in result[0].get_text()
        assert list(result[0].widgets())[0].field_value == 'preserved'


def test_validation_and_local_origin(pdf):
    source=register(pdf)
    assert client.post('/api/export',json={**payload(source),'range':'0-5'}).status_code==422
    assert client.post('/api/import',files={'file':('bad.pdf',b'invalid')}).status_code==422
    assert client.post('/api/export',json=payload(source,'<svg xmlns="http://www.w3.org/2000/svg"><image href="file:///etc/passwd"/></svg>')).status_code==422
    assert client.post('/api/export',json=payload(source),headers={'origin':'https://untrusted.example'}).status_code==403


def test_image_and_text_import():
    img=io.BytesIO(); Image.new('RGB',(400,200),'red').save(img,'PNG')
    response=client.post('/api/import',files={'file':('image.png',img.getvalue())})
    assert response.status_code==200
    assert response.json()['pages'][0]['width']==400
    response=client.post('/api/import',files={'file':('notes.txt',b'A simple note\nSecond line')})
    assert response.status_code==200, response.text
    with fitz.open(stream=base64.b64decode(response.json()['data']),filetype='pdf') as doc:
        assert 'A simple note' in doc[0].get_text()


@pytest.mark.skipif(not libreoffice(),reason='LibreOffice not installed')
def test_office_conversion():
    doc=Document(); doc.add_heading('Office conversion test',0); doc.add_paragraph('A real Word document converted locally.')
    buf=io.BytesIO(); doc.save(buf)
    response=client.post('/api/import',files={'file':('test.docx',buf.getvalue())})
    assert response.status_code==200, response.text
    with fitz.open(stream=base64.b64decode(response.json()['data']),filetype='pdf') as pdf:
        assert 'Office conversion test' in pdf[0].get_text()
